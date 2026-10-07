import io
import re
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment

from core.modules import PFN_MODULES

# -----------------------------------------------------------------------------
# Constants
# -----------------------------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parent.parent
COLUMN_DICTIONARY_PATH = PROJECT_ROOT / "data" / "metadata_column_dictionary.xlsx"
COLUMN_DICTIONARY_LABEL = COLUMN_DICTIONARY_PATH.relative_to(PROJECT_ROOT).as_posix()

# Excel sheet names
SHEET_STUDY_SETUP = "Study setup"   # study information
SHEET_LCI = "LCI"
SHEET_LCIA = "LCIA"
SHEET_METADATA = "Metadata"         # meaning of the columns (from the column dictionary file)

MAX_FILENAME_PART = 40              # keep file names short, whatever the study name length

# -----------------------------------------------------------------------------
# Helpers (functions only; cached so files are not rebuilt on every rerun)
# -----------------------------------------------------------------------------


def slugify(text, max_len=MAX_FILENAME_PART):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", str(text)).strip("_")[:max_len]


def unique_text(series):
    """Distinct non-missing values as one string (normally a single value)."""
    return ", ".join(series.dropna().astype(str).unique().tolist())


@st.cache_data
def load_column_dictionary(path, mtime):
    # `mtime` only takes part in the cache key: editing the Excel file refreshes the cache automatically.
    # sheet_name=0 reads the first sheet, so renaming the sheet inside the file does not break the app.
    df = pd.read_excel(path, sheet_name=0, dtype=str).fillna("")
    df["Column"] = df["Column"].str.strip()
    return df.drop_duplicates(subset="Column", keep="first")


def build_dictionary_sheet(dictionary_df, *tables):
    """Dictionary rows for the columns present in the exported tables (dictionary order),
    plus a placeholder row for any exported column that is not documented yet."""
    present = list(dict.fromkeys(c for t in tables for c in t.columns))
    documented = dictionary_df[dictionary_df["Column"].isin(present)]
    known = set(dictionary_df["Column"])
    undocumented = [c for c in present if c not in known]
    placeholders = pd.DataFrame({"Column": undocumented, "Meaning": "(not documented yet)"})
    return pd.concat([documented, placeholders], ignore_index=True), undocumented


@st.cache_data
def to_csv_bytes(df):
    return df.to_csv(index=False).encode("utf-8")


@st.cache_data
def to_json_str(df):
    return df.to_json(orient="records", indent=2)


@st.cache_data
def to_excel_bytes(study_setup_df, lci_df, lcia_df, metadata_df=None):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        study_setup_df.to_excel(writer, index=False, sheet_name=SHEET_STUDY_SETUP)
        lci_df.to_excel(writer, index=False, sheet_name=SHEET_LCI)
        lcia_df.to_excel(writer, index=False, sheet_name=SHEET_LCIA)
        if metadata_df is not None:
            metadata_df.to_excel(writer, index=False, sheet_name=SHEET_METADATA)
            ws = writer.sheets[SHEET_METADATA]
            ws.column_dimensions["A"].width = 36
            ws.column_dimensions["B"].width = 120
            for row in ws.iter_rows(min_row=2):
                row[1].alignment = Alignment(wrap_text=True, vertical="top")
    return buf.getvalue()


# -----------------------------------------------------------------------------
# Title
# -----------------------------------------------------------------------------

st.title(":material/download: Export")

study_id = st.session_state.current_study_id
if study_id is None:
    st.warning("Create or select a study on the :material/settings: **Study setup** page first.")
    st.stop()

study = st.session_state.studies[study_id]

# Results are stored by pages/3_Results.py, keyed by study id.
results = st.session_state.get("last_results", {}).get(study_id)
if results is None:
    st.info("No results computed yet for this study -- visit the :material/bar_chart: **Results** page first.")
    st.stop()

# Staleness guard: refuse to export results computed from a different set of flows
current_flow_ids = tuple(f["id"] for m in study["modules"] for f in study["flows"][m])
if current_flow_ids != results["flow_ids"]:
    st.warning(
        "The flows in this study have changed since the results were last computed. "
        "Visit the :material/bar_chart: **Results** page to refresh them before exporting."
    )
    st.stop()

# -----------------------------------------------------------------------------
# Prepare the data to export
# -----------------------------------------------------------------------------

lci_df = results["lci_df"]
lcia_df = results["lcia_df"]
impact_method = results["impact_method"]

study_name = study["study_name"]
indicator_level = unique_text(lcia_df["impact_indicator_level"])
temporal_representation = unique_text(lcia_df["temporal_representation"])

# Study information sheet
study_rows = [
    ("Study name", study_name),
    ("Functional unit", study["functional_unit"]),
    ("Modules", ", ".join(PFN_MODULES[m]["label"] for m in study["modules"])),
    ("Impact method", impact_method),
    ("Impact indicator level", indicator_level),
    ("Temporal representation", temporal_representation),
    ("Impact score unit", unique_text(lcia_df["unit"])),
]
if study.get("code"):
    study_rows.insert(0, ("Study code", study["code"]))
study_setup_df = pd.DataFrame(study_rows, columns=["Field", "Value"])

# Column descriptions sheet
if COLUMN_DICTIONARY_PATH.exists():
    dictionary_df = load_column_dictionary(COLUMN_DICTIONARY_PATH, COLUMN_DICTIONARY_PATH.stat().st_mtime)
    metadata_sheet_df, undocumented_columns = build_dictionary_sheet(dictionary_df, lci_df, lcia_df)
else:
    metadata_sheet_df, undocumented_columns = None, []

# File names
file_stem_lci = f"{slugify(study_name)}_lci"
file_stem_lcia = f"{slugify(study_name)}_lcia_{slugify(impact_method)}"
excel_file_name = "_".join(
    slugify(part) for part in [study_name, impact_method, indicator_level, temporal_representation] if part
) + "_results.xlsx"

# -----------------------------------------------------------------------------
# Header and data-quality notes
# -----------------------------------------------------------------------------

st.markdown(f"#### :material/label: Study: **{study_name}** | Impact method: **{impact_method}**")

# Flag emissions that could not be characterized, so they are not silently exported as empty impacts
n_missing_cf = int(lcia_df["cf_value"].isna().sum())
if n_missing_cf:
    st.warning(
        f"{n_missing_cf} emission row(s) have no matching characterization factor for "
        f"'{impact_method}' and therefore have an empty `impact_score` in the LCIA table."
    )

if metadata_sheet_df is None:
    st.caption(f":material/warning: Column dictionary not found (`{COLUMN_DICTIONARY_LABEL}`); "
               f"the Excel workbook will not include the *{SHEET_METADATA}* sheet.")
elif undocumented_columns:
    st.caption(
        f":material/warning: {len(undocumented_columns)} column(s) are not described yet in "
        f"`{COLUMN_DICTIONARY_LABEL}`: {', '.join(undocumented_columns)}"
    )

# -----------------------------------------------------------------------------
# Preview
# -----------------------------------------------------------------------------

st.subheader(":material/table_view: Results tables")

table_choice = st.radio(
    "Table",
    ["lci", "lcia"],
    format_func=lambda t: {
        "lci": "Microplastic emissions (LCI)",
        "lcia": "Potential impacts (LCIA)",
    }[t],
    horizontal=True,
)

if table_choice == "lcia":
    df_export, file_stem = lcia_df, file_stem_lcia
else:
    df_export, file_stem = lci_df, file_stem_lci

PREVIEW_ROWS = 5
st.dataframe(df_export.head(PREVIEW_ROWS), width="stretch")
st.caption(f"Preview of the first {min(PREVIEW_ROWS, len(df_export))} of {len(df_export)} rows. "
           "The downloads contain all rows.")

# -----------------------------------------------------------------------------
# Downloads
# -----------------------------------------------------------------------------

st.subheader(":material/download: Downloads")
st.caption(
    "CSV and JSON contain the table selected above. "
    f"The Excel workbook contains all tables ({SHEET_STUDY_SETUP}, {SHEET_LCI}, {SHEET_LCIA}) "
    f"and a {SHEET_METADATA} sheet describing the columns."
)

c1, c2, c3 = st.columns(3)
with c1:
    st.download_button(
        ":material/csv: Download CSV", data=to_csv_bytes(df_export),
        file_name=f"{file_stem}.csv", mime="text/csv",
    )
with c2:
    st.download_button(
        ":material/table: Download Excel workbook (all tables)",
        data=to_excel_bytes(study_setup_df, lci_df, lcia_df, metadata_sheet_df),
        file_name=excel_file_name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
with c3:
    st.download_button(
        ":material/data_object: Download JSON", data=to_json_str(df_export),
        file_name=f"{file_stem}.json", mime="application/json",
    )

# -----------------------------------------------------------------------------
# Traceability
# -----------------------------------------------------------------------------

st.divider()
st.subheader(":material/fact_check: Traceability")
st.markdown(
    f"- **Impact method and CFs**: the method is recorded in the Excel *{SHEET_STUDY_SETUP}* sheet and in the "
    "`impact_method` column; CF values and their 95 % bounds are in `cf_value`, `cf_ll_95` and `cf_ul_95`.\n"
    "- **Flow parameters**: for flows estimated from macroplastics, each row also carries the "
    "intermediate parameters applied (e.g. release rates, fragmentation rates, loss rates).\n"
    "- **Defaults vs. user overrides**: not tracked -- the exported values are the final values "
    "used, whether kept as defaults or edited.\n"
    "- **Calculation**: `micro_mass_kg` (LCI) is multiplied by `cf_value` to obtain `impact_score` (LCIA).\n"
    f"- **Column descriptions**: the meaning of each column is given in the Excel *{SHEET_METADATA}* sheet."
)